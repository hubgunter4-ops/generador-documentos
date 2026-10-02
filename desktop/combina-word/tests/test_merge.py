from pathlib import Path

from docx import Document

from app.merge import (
    detect_template_fields,
    find_unresolved_placeholders,
    merge_excel_into_one_docx,
    replace_in_document,
    validate_template_fields,
)


def test_replaces_paragraph_and_table(tmp_path: Path):
    path = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Nombre: {{NOMBRE}}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "ID: {{ID}}"
    doc.save(path)

    rendered = Document(path)
    replace_in_document(rendered, {"NOMBRE": "Ana", "ID": "A-1"})

    assert "Ana" in rendered.paragraphs[0].text
    assert "A-1" in rendered.tables[0].cell(0, 0).text
    assert find_unresolved_placeholders(rendered) == []


def test_excel_headers_match_lowercase_word_markers(tmp_path: Path):
    import pandas as pd

    template = tmp_path / "template.docx"
    excel = tmp_path / "data.xlsx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("Yo: {{yo}}")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Padres: {{mis_padres}}"
    doc.save(template)
    pd.DataFrame([{"yo": "Ana", "mis_padres": "Luis / Marta"}]).to_excel(excel, index=False)

    assert merge_excel_into_one_docx(excel, template, output) == 1
    rendered = Document(output)
    assert "Ana" in rendered.paragraphs[0].text
    assert "Luis / Marta" in rendered.tables[0].cell(0, 0).text


def test_record_limit_generates_requested_copy_count(tmp_path: Path):
    import pandas as pd

    template = tmp_path / "template.docx"
    excel = tmp_path / "data.xlsx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("Yo: {{yo}}")
    doc.save(template)
    pd.DataFrame([{"yo": "Ana"}, {"yo": "Luis"}, {"yo": "Marta"}]).to_excel(excel, index=False)

    assert merge_excel_into_one_docx(excel, template, output, record_limit=2) == 2
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "Ana" in text and "Luis" in text and "Marta" not in text


def test_detects_and_fills_visible_categories(tmp_path: Path):
    import pandas as pd

    template = tmp_path / "category-template.docx"
    excel = tmp_path / "category-data.xlsx"
    output = tmp_path / "category-output.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "mis_padres"
    table.cell(0, 1).text = "yo"
    table.cell(1, 0).text = "Nombre(s):\n____________________________"
    table.cell(1, 1).text = "Nombre(s):\n____________________________"
    doc.save(template)
    pd.DataFrame([{"mis_padres": "Luis / Marta", "yo": "Ana"}]).to_excel(excel, index=False)

    assert detect_template_fields(template) == ["MIS_PADRES", "YO"]
    assert validate_template_fields(template, excel)[1] == []
    assert merge_excel_into_one_docx(excel, template, output) == 1
    rendered = Document(output)
    text = "\n".join(cell.text for row in rendered.tables[0].rows for cell in row.cells)
    assert "Luis / Marta" in text and "Ana" in text
