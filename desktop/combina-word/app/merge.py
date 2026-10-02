from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
import re

import pandas as pd
from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer


def format_value(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_headers(columns) -> list[str]:
    return [str(column).strip().upper() for column in columns]


KNOWN_CATEGORY_FIELDS = {
    "MIS_PADRES",
    "MIS_PADRINOS",
    "ABUELOS_PATERNOS",
    "ABUELOS_MATERNOS",
    "EL_SR_Y_LA_SRA",
    "YO",
}

INFERRED_LABELS = (
    (r"\bYO\b", "YO"),
    (r"\bHOY\b", "FECHA_BAUTIZO"),
    (r"\bNAC[IÍ]\s+EL\b", "FECHA_NACIMIENTO"),
    (r"\bEN\s+SANTA\s+CRUZ\b|\bLUGAR\s+DE\s+NACIMIENTO\b", "LUGAR_NACIMIENTO"),
    (r"\bPAP[AÁ]\b", "PAPA"),
    (r"\bMAM[AÁ]\b", "MAMA"),
    (r"\bABUELOS\s+PATERNOS\b", "ABUELOS_PATERNOS"),
    (r"\bABUELOS\s+MATERNOS\b", "ABUELOS_MATERNOS"),
    (r"\bMIS\s+PADRINOS\b", "MIS_PADRINOS"),
    (r"\bLIBRO\b", "LIBRO"),
    (r"\bP[AÁ]G\.?\b", "PAGINA"),
    (r"\bPART\.?\b", "PARTIDA"),
    (r"\bEXPEDIDA\b", "FECHA_EXPEDIDA"),
)


def normalize_field(value: str) -> str:
    return str(value).strip().strip("{}").strip().upper()


def _table_paragraphs(tables):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                yield from _table_paragraphs(cell.tables)


def _document_paragraphs(document: Document):
    yield from document.paragraphs
    yield from _table_paragraphs(document.tables)
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from container.paragraphs
            yield from _table_paragraphs(container.tables)


def detect_template_fields(template_path: str | Path) -> list[str]:
    """Detect {{FIELD}} markers and visible category headings in a DOCX."""
    document = Document(str(template_path))
    fields: set[str] = set()
    for paragraph in _document_paragraphs(document):
        fields.update(normalize_field(match) for match in re.findall(r"\{\{([^{}]+)\}\}", paragraph.text))
        if re.search(r"_{3,}|\.{3,}", paragraph.text):
            for pattern, field in INFERRED_LABELS:
                if re.search(pattern, paragraph.text, re.IGNORECASE):
                    fields.add(field)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                candidate = normalize_field(cell.text)
                if candidate in KNOWN_CATEGORY_FIELDS:
                    fields.add(candidate)
    return sorted(field for field in fields if field)


def replace_in_paragraph(paragraph, data: dict[str, Any]) -> None:
    """Replace {{FIELD}} placeholders, including placeholders split across runs."""
    if not paragraph.runs:
        return
    text = "".join(run.text for run in paragraph.runs)
    for field, value in data.items():
        marker = re.compile(r"\{\{" + re.escape(str(field)) + r"\}\}", re.IGNORECASE)
        text = marker.sub(lambda _match: format_value(value), text)
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_in_tables(tables, data: dict[str, Any]) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, data)
                replace_in_tables(cell.tables, data)


def fill_visible_category_tables(document: Document, data: dict[str, Any]) -> None:
    """Fill the capture cell below visible category headings without {{markers}}."""
    for table in document.tables:
        for header_index, row in enumerate(table.rows[:-1]):
            fields = [normalize_field(cell.text) for cell in row.cells]
            if not any(field in data for field in fields):
                continue
            target_row = table.rows[header_index + 1]
            for index, field in enumerate(fields):
                if field not in data or index >= len(target_row.cells):
                    continue
                target_cell = target_row.cells[index]
                for paragraph in target_cell.paragraphs:
                    if re.search(r"_{3,}", paragraph.text):
                        paragraph.text = re.sub(r"_{3,}", format_value(data[field]), paragraph.text, count=1)
                        break
                for paragraph in target_cell.paragraphs:
                    if re.fullmatch(r"\s*_{3,}\s*", paragraph.text):
                        paragraph.text = ""


def fill_inferred_blank_paragraphs(document: Document, data: dict[str, Any]) -> None:
    """Fill underline/dot blanks when a recognizable label precedes the blank."""
    for paragraph in _document_paragraphs(document):
        if not re.search(r"_{3,}|\.{3,}", paragraph.text):
            continue
        matches = []
        for pattern, field in INFERRED_LABELS:
            if field not in data:
                continue
            match = re.search(pattern, paragraph.text, re.IGNORECASE)
            if match:
                matches.append((match.start(), match.end(), field))
        for start, end, field in sorted(matches, reverse=True):
            next_starts = [item[0] for item in matches if item[0] > start]
            segment_end = min(next_starts) if next_starts else len(paragraph.text)
            segment = paragraph.text[end:segment_end]
            segment = re.sub(r"_{3,}|\.{3,}", format_value(data[field]), segment, count=1)
            paragraph.text = paragraph.text[:end] + segment + paragraph.text[segment_end:]


def replace_in_document(document: Document, data: dict[str, Any]) -> None:
    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph, data)
    replace_in_tables(document.tables, data)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                replace_in_paragraph(paragraph, data)
            replace_in_tables(container.tables, data)
    fill_visible_category_tables(document, data)
    fill_inferred_blank_paragraphs(document, data)


def add_page_break_at_end(document: Document) -> None:
    """Make the next appended Excel record begin on a new physical page."""
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def find_unresolved_placeholders(document: Document) -> list[str]:
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    for section in document.sections:
        for container in (section.header, section.footer):
            texts.extend(p.text for p in container.paragraphs)
    return sorted(set(re.findall(r"\{\{[^{}]+\}\}", "\n".join(texts))))


def load_records(excel_path: str | Path) -> pd.DataFrame:
    frame = pd.read_excel(excel_path)
    if frame.empty:
        raise ValueError("El archivo Excel no contiene registros.")
    frame.columns = normalize_headers(frame.columns)
    return frame


def validate_template_fields(template_path: str | Path, excel_path: str | Path) -> tuple[list[str], list[str], list[str]]:
    template_fields = detect_template_fields(template_path)
    excel_fields = set(load_records(excel_path).columns)
    matched = sorted(set(template_fields) & excel_fields)
    missing = sorted(set(template_fields) - excel_fields)
    unused = sorted(excel_fields - set(template_fields))
    return matched, missing, unused


def merge_excel_into_one_docx(
    excel_path: str | Path,
    template_path: str | Path,
    output_path: str | Path,
    progress_callback=None,
    record_limit: int | None = None,
) -> int:
    """Create one complete template copy per Excel row and append all copies."""
    records = load_records(excel_path)
    _matched, missing, _unused = validate_template_fields(template_path, excel_path)
    if missing:
        raise ValueError("Faltan en Excel las categorías de la plantilla: " + ", ".join(missing))
    if record_limit is not None:
        if record_limit < 1:
            raise ValueError("La cantidad de copias debe ser mayor que cero.")
        if record_limit > len(records):
            raise ValueError(f"El Excel solo contiene {len(records)} registros; no se pueden generar {record_limit} copias.")
        records = records.iloc[:record_limit].copy()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with TemporaryDirectory(prefix="combina-word-") as temp_dir:
        temp = Path(temp_dir)
        generated: list[Path] = []
        total = len(records)
        for index, (_, row) in enumerate(records.iterrows(), start=1):
            document = Document(str(template_path))
            replace_in_document(document, row.to_dict())
            if index < total:
                add_page_break_at_end(document)
            item = temp / f"registro_{index}.docx"
            document.save(item)
            generated.append(item)
            if progress_callback:
                progress_callback(index, total)

        final = Document(str(generated[0]))
        composer = Composer(final)
        for item in generated[1:]:
            composer.append(Document(str(item)))
        composer.save(str(output_path))
    return len(records)
